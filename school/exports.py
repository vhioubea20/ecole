from io import BytesIO

from django.contrib.auth import get_user_model
from django.db.models import Count
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import SchoolClass, SchoolFee, Student, StudentPayment, UserProfile

User = get_user_model()


def _safe_value(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_safe_value(value) for value in row])

    header_fill = PatternFill(fill_type="solid", fgColor="1A486C")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, cells in enumerate(sheet.columns, start=1):
        width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 42)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)


def build_school_workbook():
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Export de gestion scolaire"

    students = Student.objects.select_related("school_class").order_by("matricule")
    _add_sheet(
        workbook,
        "Élèves",
        [
            "Matricule", "Prénom", "Nom", "Date de naissance", "Classe", "Adresse", "Photo",
            "Personne responsable", "Numéro responsable", "Profession responsable", "Statut", "Date inscription",
        ],
        (
            [
                student.matricule,
                student.first_name,
                student.last_name,
                student.date_of_birth,
                student.school_class.libelle,
                student.address,
                student.photo.name,
                student.guardian_name,
                student.guardian_phone,
                student.guardian_profession,
                student.get_status_display(),
                student.enrollment_date,
            ]
            for student in students
        ),
    )

    classes = SchoolClass.objects.annotate(student_count=Count("students")).order_by("libelle")
    _add_sheet(
        workbook,
        "Classes",
        ["ID", "Libellé", "Nombre d'élèves"],
        ([school_class.pk, school_class.libelle, school_class.student_count] for school_class in classes),
    )

    fees = SchoolFee.objects.select_related("school_class").order_by("-annee", "school_class__libelle", "nature")
    _add_sheet(
        workbook,
        "Frais",
        ["Classe", "Année scolaire", "Nature", "Mode", "Nombre d'échéances", "Montant saisi (GNE)", "Montant par tranche (GNE)", "Dernière tranche (GNE)"],
        (
            [
                fee.school_class.libelle,
                fee.annee,
                fee.get_nature_display(),
                fee.get_mode_display(),
                fee.nombre_echeances,
                fee.montant,
                fee.montant_par_tranche,
                fee.montant_derniere_tranche,
            ]
            for fee in fees
        ),
    )

    payments = StudentPayment.objects.select_related("student", "student__school_class")
    _add_sheet(
        workbook,
        "Paiements",
        ["Date", "Matricule", "Prénom", "Nom", "Classe", "Montant (GNE)"],
        (
            [
                payment.date_paiement,
                payment.student.matricule,
                payment.student.first_name,
                payment.student.last_name,
                payment.student.school_class.libelle,
                payment.montant,
            ]
            for payment in payments
        ),
    )

    users = User.objects.select_related("profile").prefetch_related("user_permissions", "groups").order_by("username")
    _add_sheet(
        workbook,
        "Utilisateurs",
        ["Identifiant", "Prénom", "Nom", "Courriel", "Rôle", "Actif", "Groupes", "Permissions"],
        (
            [
                user.username,
                user.first_name,
                user.last_name,
                user.email,
                "Admin" if user.is_staff else "Utilisateur",
                "Oui" if user.is_active else "Non",
                ", ".join(group.name for group in user.groups.all()),
                ", ".join(sorted(user.get_all_permissions())),
            ]
            for user in users
        ),
    )

    profiles = UserProfile.objects.select_related("user").order_by("user__username")
    _add_sheet(
        workbook,
        "Profils",
        ["Identifiant", "Photo"],
        ([profile.user.username, profile.photo.name] for profile in profiles),
    )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()