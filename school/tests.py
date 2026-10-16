from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .forms import SchoolClassForm, SchoolFeeForm, StudentForm, StudentPaymentForm, UserCreateForm, UserUpdateForm
from .models import SchoolClass, SchoolFee, Student, StudentPayment, UserProfile
from .services import school_year_data
from django.contrib.auth.models import Permission

User = get_user_model()


class SchoolAccessTests(TestCase):
    def setUp(self):
        self.staff = User.objects.create_user(username="staff", password="Strong-pass-2026", is_staff=True)
        self.member = User.objects.create_user(username="member", password="Strong-pass-2026")

    def test_dashboard_requires_authentication(self):
        response = self.client.get(reverse("dashboard"))
        self.assertRedirects(response, f"{reverse('login')}?next={reverse('dashboard')}")

    def test_staff_can_export_all_tables_without_password_data(self):
        school_class = SchoolClass.objects.create(libelle="CP")
        student = Student.objects.create(first_name="Awa", last_name="Diallo", school_class=school_class)
        year = SchoolFeeForm().fields["annee"].choices[0][0]
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year,
            nature=SchoolFee.Nature.TUITION,
            montant="90000.00",
        )
        StudentPayment.objects.create(student=student, date_paiement=date(2026, 10, 15), montant="10000.00")
        self.client.force_login(self.staff)

        response = self.client.get(reverse("export_all_excel"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Élèves", "Classes", "Frais", "Paiements", "Utilisateurs", "Profils"])
        user_headers = next(workbook["Utilisateurs"].iter_rows(values_only=True))
        self.assertNotIn("Mot de passe", user_headers)
        self.assertNotIn("password", user_headers)
        self.assertEqual(workbook["Élèves"].max_row, 2)

    def test_regular_user_cannot_export_all_tables(self):
        self.client.force_login(self.member)

        response = self.client.get(reverse("export_all_excel"))

        self.assertEqual(response.status_code, 302)

    def test_login_page_loads_password_visibility_control(self):
        response = self.client.get(reverse("login"))
        self.assertContains(response, "js/password-toggle.js")

        self.client.force_login(self.staff)
        response = self.client.get(reverse("user_create"))
        self.assertContains(response, "js/password-toggle.js")
        self.assertContains(response, 'name="password1"')
        self.assertContains(response, 'name="password2"')

    def test_staff_can_open_student_registration(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("student_create"))
        self.assertEqual(response.status_code, 200)

    def test_student_badge_contains_requested_details(self):
        school_class = SchoolClass.objects.create(libelle="CM1")
        student = Student.objects.create(
            first_name="Awa",
            last_name="Diallo",
            date_of_birth=date(2015, 3, 12),
            school_class=school_class,
            address="12 rue des Écoles",
            guardian_name="Mariam Diallo",
            guardian_phone="0102030405",
        )
        self.client.force_login(self.staff)

        response = self.client.get(reverse("student_badge", args=[student.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, student.matricule)
        self.assertContains(response, "Awa Diallo")
        self.assertContains(response, "12/03/2015")
        self.assertContains(response, "CM1")
        self.assertContains(response, "12 rue des Écoles")
        self.assertContains(response, "Mariam Diallo")
        self.assertContains(response, "0102030405")
        self.assertContains(response, "Signature du Directeur")

    def test_regular_user_cannot_manage_users(self):
        self.client.force_login(self.member)
        response = self.client.get(reverse("user_list"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response.url)

    def test_admin_role_grants_all_rights_without_email(self):
        form = UserCreateForm(
            data={
                "username": "newadmin",
                "first_name": "Admin",
                "last_name": "Test",
                "email": "",
                "role": "admin",
                "password1": "Good-Test-pass-2026!",
                "password2": "Good-Test-pass-2026!",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)

        user = form.save()

        self.assertFalse(user.is_superuser)
        self.assertTrue(user.is_staff)
        self.assertEqual(user.email, "")
        self.assertTrue(user.has_perm("school.delete_student"))
        self.assertTrue(UserProfile.objects.filter(user=user).exists())

    def test_regular_user_receives_only_selected_permissions(self):
        permission = Permission.objects.get(content_type__app_label="school", codename="view_student")
        form = UserCreateForm(
            data={
                "username": "reader",
                "first_name": "Read",
                "last_name": "Only",
                "email": "",
                "role": "user",
                "permissions": [permission.pk],
                "password1": "Good-Test-pass-2026!",
                "password2": "Good-Test-pass-2026!",
            }
        )
        self.assertTrue(form.is_valid(), form.errors)
        user = form.save()
        self.assertFalse(user.is_staff)
        self.assertTrue(user.has_perm("school.view_student"))
        self.assertFalse(user.has_perm("school.add_student"))

        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse("student_list")).status_code, 200)
        self.assertEqual(self.client.get(reverse("class_list")).status_code, 302)

    def test_downgrading_admin_does_not_keep_all_permissions_by_default(self):
        admin = User.objects.create_user(username="appadmin", password="Strong-pass-2026", is_staff=True)
        admin.user_permissions.set(Permission.objects.filter(content_type__app_label="school"))
        form = UserUpdateForm(instance=admin)

        self.assertFalse(form["permissions"].value())

    def test_student_form_requires_only_first_name_last_name_and_class(self):
        school_class = SchoolClass.objects.create(libelle="CM1")
        form = StudentForm(data={"first_name": "Jean Pierre", "last_name": "Dupont", "school_class": school_class.pk})
        self.assertTrue(form.is_valid(), form.errors)

    def test_matricule_uses_name_prefix_and_database_id(self):
        school_class = SchoolClass.objects.create(libelle="CM1")
        student = Student.objects.create(first_name="Jean Pierre", last_name="Dupont", school_class=school_class)
        self.assertEqual(student.matricule, f"JDU{student.pk:04d}")

    def test_staff_can_open_class_management(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("class_list"))
        self.assertEqual(response.status_code, 200)

    def test_class_label_is_required(self):
        self.assertFalse(SchoolClassForm(data={"libelle": ""}).is_valid())

    def test_staff_can_create_and_update_class(self):
        self.client.force_login(self.staff)
        response = self.client.post(reverse("class_create"), {"libelle": "CP"})
        self.assertRedirects(response, reverse("class_list"))
        school_class = SchoolClass.objects.get(libelle="CP")

        response = self.client.post(
            reverse("class_update", args=[school_class.pk]),
            {"libelle": "CP A"},
        )
        self.assertRedirects(response, reverse("class_list"))
        school_class.refresh_from_db()
        self.assertEqual(school_class.libelle, "CP A")

    def test_class_with_students_cannot_be_deleted(self):
        school_class = SchoolClass.objects.create(libelle="CM1")
        Student.objects.create(first_name="Awa", last_name="Diallo", school_class=school_class)
        self.client.force_login(self.staff)

        response = self.client.post(reverse("class_delete", args=[school_class.pk]))

        self.assertRedirects(response, reverse("class_list"))
        self.assertTrue(SchoolClass.objects.filter(pk=school_class.pk).exists())

    def test_fee_requires_year_class_and_amount(self):
        self.assertFalse(SchoolFeeForm(data={}).is_valid())

    def test_same_class_cannot_have_two_fees_of_same_nature_for_same_year(self):
        school_class = SchoolClass.objects.create(libelle="CP")
        year = SchoolFeeForm().fields["annee"].choices[0][0]
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year,
            nature=SchoolFee.Nature.TUITION,
            montant="100000.00",
        )

        form = SchoolFeeForm(
            data={
                "school_class": school_class.pk,
                "annee": year,
                "nature": SchoolFee.Nature.TUITION,
                "montant": "120000.00",
            }
        )

        self.assertFalse(form.is_valid())

    def test_database_rejects_duplicate_class_year_nature_fee(self):
        school_class = SchoolClass.objects.create(libelle="CP")
        year = SchoolFeeForm().fields["annee"].choices[0][0]
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year,
            nature=SchoolFee.Nature.REGISTRATION,
            montant="100000.00",
        )

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                SchoolFee.objects.create(
                    school_class=school_class,
                    annee=year,
                    nature=SchoolFee.Nature.REGISTRATION,
                    montant="120000.00",
                )

    def test_same_class_and_year_allow_one_fee_of_each_nature(self):
        school_class = SchoolClass.objects.create(libelle="CP")
        year = SchoolFeeForm().fields["annee"].choices[0][0]
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year,
            nature=SchoolFee.Nature.REGISTRATION,
            montant="100000.00",
        )

        form = SchoolFeeForm(
            data={
                "school_class": school_class.pk,
                "annee": year,
                "nature": SchoolFee.Nature.TUITION,
                "montant": "120000.00",
            }
        )

        self.assertTrue(form.is_valid(), form.errors)
        form.save()
        self.assertEqual(SchoolFee.objects.filter(school_class=school_class, annee=year).count(), 2)

    def test_staff_can_open_fee_management(self):
        self.client.force_login(self.staff)
        response = self.client.get(reverse("fee_list"))
        self.assertEqual(response.status_code, 200)

    def test_staff_can_create_update_and_delete_fee(self):
        school_class = SchoolClass.objects.create(libelle="CP")
        year = SchoolFeeForm().fields["annee"].choices[0][0]
        self.client.force_login(self.staff)

        response = self.client.post(
            reverse("fee_create"),
            {
                "school_class": school_class.pk,
                "annee": year,
                "nature": SchoolFee.Nature.TUITION,
                "montant": "100000.00",
            },
        )
        self.assertRedirects(response, reverse("fee_list"))
        fee = SchoolFee.objects.get(school_class=school_class, annee=year)

        response = self.client.post(
            reverse("fee_update", args=[fee.pk]),
            {
                "school_class": school_class.pk,
                "annee": year,
                "nature": SchoolFee.Nature.TUITION,
                "montant": "125000.00",
            },
        )
        self.assertRedirects(response, reverse("fee_list"))
        fee.refresh_from_db()
        self.assertEqual(str(fee.montant), "125000.00")

        response = self.client.post(reverse("fee_delete", args=[fee.pk]))
        self.assertRedirects(response, reverse("fee_list"))
        self.assertFalse(SchoolFee.objects.filter(pk=fee.pk).exists())

    def test_class_with_fees_cannot_be_deleted(self):
        school_class = SchoolClass.objects.create(libelle="CE1")
        year = SchoolFeeForm().fields["annee"].choices[0][0]
        SchoolFee.objects.create(school_class=school_class, annee=year, montant="100000.00")
        self.client.force_login(self.staff)

        response = self.client.post(reverse("class_delete", args=[school_class.pk]))

        self.assertRedirects(response, reverse("class_list"))
        self.assertTrue(SchoolClass.objects.filter(pk=school_class.pk).exists())

    def test_payment_form_uses_date_matricule_and_amount(self):
        school_class = SchoolClass.objects.create(libelle="CM1")
        student = Student.objects.create(first_name="Awa", last_name="Diallo", school_class=school_class)
        form = StudentPaymentForm(
            data={"date_paiement": "2026-10-15", "student": student.pk, "montant": "25000.00"}
        )

        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.fields["student"].label, "Matricule élève")

    def test_schooling_overview_sums_due_and_year_payments(self):
        year = school_year_data()
        school_class = SchoolClass.objects.create(libelle="CM1")
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year["school_year"],
            nature=SchoolFee.Nature.REGISTRATION,
            montant="10000.00",
        )
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year["school_year"],
            nature=SchoolFee.Nature.TUITION,
            montant="90000.00",
        )
        first_student = Student.objects.create(first_name="Awa", last_name="Diallo", school_class=school_class)
        Student.objects.create(first_name="Moussa", last_name="Ba", school_class=school_class)
        StudentPayment.objects.create(
            student=first_student,
            date_paiement=year["months"][0]["start"],
            montant="10000.00",
        )
        StudentPayment.objects.create(
            student=first_student,
            date_paiement=year["period_end"],
            montant="5000.00",
        )
        self.client.force_login(self.staff)

        response = self.client.get(reverse("schooling_overview"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_due"], Decimal("200000.00"))
        self.assertEqual(response.context["total_paid"], Decimal("10000.00"))

    def test_class_and_student_details_show_monthly_and_payment_history(self):
        year = school_year_data()
        school_class = SchoolClass.objects.create(libelle="CP")
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year["school_year"],
            nature=SchoolFee.Nature.REGISTRATION,
            montant="10000.00",
        )
        SchoolFee.objects.create(
            school_class=school_class,
            annee=year["school_year"],
            nature=SchoolFee.Nature.TUITION,
            montant="90000.00",
        )
        student = Student.objects.create(first_name="Jean", last_name="Kone", school_class=school_class)
        StudentPayment.objects.create(
            student=student,
            date_paiement=year["months"][0]["start"],
            montant="10000.00",
        )
        self.client.force_login(self.staff)

        class_response = self.client.get(reverse("schooling_class_detail", args=[school_class.pk]))
        student_response = self.client.get(reverse("student_payment_detail", args=[student.pk]))

        self.assertEqual(len(class_response.context["months"]), 9)
        self.assertEqual(class_response.context["students"][0]["monthly_amounts"][0], Decimal("10000.00"))
        self.assertEqual(class_response.context["annual_fee"], Decimal("100000.00"))
        self.assertEqual(student_response.context["total_paid"], Decimal("10000.00"))
        self.assertEqual(student_response.context["amount_due"], Decimal("100000.00"))
        self.assertEqual(len(student_response.context["payments"]), 1)

    def test_staff_can_create_update_and_delete_student_payment(self):
        school_class = SchoolClass.objects.create(libelle="CE1")
        student = Student.objects.create(first_name="Ali", last_name="Traore", school_class=school_class)
        self.client.force_login(self.staff)
        payload = {"date_paiement": "2026-10-15", "student": student.pk, "montant": "10000.00"}

        response = self.client.post(reverse("student_payment_create"), payload)
        self.assertRedirects(response, reverse("student_payment_detail", args=[student.pk]))
        payment = StudentPayment.objects.get(student=student)

        response = self.client.post(
            reverse("student_payment_update", args=[payment.pk]),
            {**payload, "montant": "12000.00"},
        )
        self.assertRedirects(response, reverse("student_payment_detail", args=[student.pk]))
        payment.refresh_from_db()
        self.assertEqual(payment.montant, Decimal("12000.00"))

        response = self.client.post(reverse("student_payment_delete", args=[payment.pk]))
        self.assertRedirects(response, reverse("student_payment_detail", args=[student.pk]))
        self.assertFalse(StudentPayment.objects.filter(pk=payment.pk).exists())

    def test_staff_can_download_payment_receipt_as_pdf(self):
        school_class = SchoolClass.objects.create(libelle="CP")
        student = Student.objects.create(first_name="Awa", last_name="Diallo", school_class=school_class)
        payment = StudentPayment.objects.create(
            student=student,
            date_paiement=date(2026, 10, 16),
            montant="25000.00",
        )
        self.client.force_login(self.staff)

        response = self.client.get(reverse("student_payment_receipt", args=[payment.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertIn(f"recu-paiement-{payment.pk:06d}.pdf", response["Content-Disposition"])

    def test_user_without_payment_permission_cannot_download_receipt(self):
        school_class = SchoolClass.objects.create(libelle="CP")
        student = Student.objects.create(first_name="Ali", last_name="Traore", school_class=school_class)
        payment = StudentPayment.objects.create(
            student=student,
            date_paiement=date(2026, 10, 16),
            montant="10000.00",
        )
        self.client.force_login(self.member)

        response = self.client.get(reverse("student_payment_receipt", args=[payment.pk]))

        self.assertEqual(response.status_code, 302)
